"""Relatório policial em DOCX e PDF."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from config import REPORT_DIR
from modules import db
from modules.correlation import graph, multi_account_imei_alerts
from modules.utils import display_dt, from_json, now_iso


def _case_pack(case_id: int) -> dict:
    case = db.get_case(case_id) or {}
    return {
        "case": case,
        "accounts": db.query("SELECT * FROM accounts WHERE case_id = ?", (case_id,)),
        "devices": db.query("SELECT * FROM devices WHERE case_id = ?", (case_id,)),
        "links": db.query("SELECT * FROM device_accounts WHERE case_id = ?", (case_id,)),
        "events": db.query(
            "SELECT * FROM events WHERE case_id = ? ORDER BY ts IS NULL, ts LIMIT 400",
            (case_id,),
        ),
        "locations": db.query(
            "SELECT * FROM locations WHERE case_id = ? ORDER BY ts IS NULL, ts LIMIT 400",
            (case_id,),
        ),
        "photos": db.query("SELECT * FROM photos WHERE case_id = ? ORDER BY taken_at LIMIT 200", (case_id,)),
        "files": db.query("SELECT * FROM files WHERE case_id = ?", (case_id,)),
        "payments": db.query("SELECT * FROM payments WHERE case_id = ?", (case_id,)),
        "imports": db.query("SELECT * FROM imports WHERE case_id = ?", (case_id,)),
        "alerts": multi_account_imei_alerts(case_id),
        "graph": graph(case_id),
    }


def _emails_of_device(pack, device_id: int) -> list[str]:
    return sorted({r["account_email"] for r in pack["links"] if r["device_id"] == device_id})


def generate_docx(case_id: int) -> Path:
    pack = _case_pack(case_id)
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    dest = REPORT_DIR / f"relatorio_forense_caso_{case_id}_{now_iso().replace(':', '')}.docx"
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    title = doc.add_paragraph()
    run = title.add_run("POLÍCIA CIVIL DO PARANÁ")
    run.bold = True
    run.font.size = Pt(14)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph()
    r2 = sub.add_run("ANALISADOR GOOGLE FORENSE — RELATÓRIO POLICIAL")
    r2.bold = True
    r2.font.size = Pt(12)
    r2.font.color.rgb = RGBColor(0x18, 0x3A, 0x31)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        "Documento gerado automaticamente a partir de produção Google LERS. "
        "Conferir sempre com os arquivos originais da resposta do provedor. "
        "Hashes SHA-256 dos ZIPs importados constam neste relatório."
    )
    doc.add_heading("1. Identificação da conta", level=1)
    if not pack["accounts"]:
        doc.add_paragraph("Nenhuma conta extraída.")
    for acc in pack["accounts"]:
        doc.add_paragraph(
            f"Conta: {acc.get('google_account') or '-'}\n"
            f"Nome: {acc.get('display_name') or '-'}\n"
            f"E-mail principal: {acc.get('primary_email') or '-'}\n"
            f"E-mails alternativos: {', '.join(from_json(acc.get('alternate_emails'), []) or []) or '-'}\n"
            f"Telefones: {', '.join(from_json(acc.get('phones'), []) or []) or '-'}\n"
            f"Criação: {display_dt(acc.get('created_on')) or '-'}\n"
            f"Última atividade: {display_dt(acc.get('last_activity')) or '-'}\n"
            f"Status: {acc.get('status') or '-'}\n"
            f"Data de exclusão: {display_dt(acc.get('deletion_date')) or '-'}"
        )
    doc.add_heading("2. Dispositivos encontrados", level=1)
    for device in pack["devices"]:
        emails = _emails_of_device(pack, device["id"])
        doc.add_paragraph(
            f"Modelo: {device.get('model') or '-'} | Fabricante: {device.get('manufacturer') or '-'}\n"
            f"IMEI 1: {device.get('imei1') or '-'} | IMEI 2: {device.get('imei2') or '-'}\n"
            f"Série: {device.get('serial_number') or '-'} | Android ID: {device.get('android_id') or '-'}\n"
            f"Primeiro vínculo: {display_dt(device.get('first_seen')) or '-'}\n"
            f"Último vínculo: {display_dt(device.get('last_seen')) or '-'}\n"
            f"Contas associadas: {', '.join(emails) or '-'}"
        )
    if pack["alerts"]:
        doc.add_heading("Alerta: mesmo aparelho vinculado a múltiplas contas", level=2)
        for alert in pack["alerts"]:
            doc.add_paragraph(f"IMEI {alert['imei']}: {', '.join(alert['accounts'])}")
    doc.add_heading("3. Linha temporal", level=1)
    for ev in pack["events"][:120]:
        doc.add_paragraph(
            f"{display_dt(ev.get('ts')) or '-'} | {ev.get('product') or '-'} | "
            f"{ev.get('event_type') or '-'} | {ev.get('description') or '-'}",
            style="List Bullet",
        )
    doc.add_heading("4. Locais relevantes", level=1)
    for loc in pack["locations"][:80]:
        permanencia = f"{int(loc['duration_seconds']//60)} min" if loc.get("duration_seconds") else "-"
        doc.add_paragraph(
            f"{display_dt(loc.get('ts')) or '-'} | {loc.get('lat')}, {loc.get('lon')} | "
            f"{loc.get('place_name') or '-'} | permanência: {permanencia}"
        )
    doc.add_heading("5. Arquivos importantes", level=1)
    important = [f for f in pack["files"] if from_json(f.get("keywords_hit"), [])]
    if not important:
        doc.add_paragraph("Nenhum arquivo com palavras-chave investigativas (PIX, CPF, BANCO, ARMA, DROGA, NOME, TELEFONE, VALOR).")
    for item in important:
        doc.add_paragraph(
            f"{item.get('name')} | {item.get('mime_type')} | palavras: "
            f"{', '.join(from_json(item.get('keywords_hit'), []) or [])}"
        )
    doc.add_heading("6. Fotos", level=1)
    for photo in pack["photos"][:60]:
        gps = f"{photo.get('lat')}, {photo.get('lon')}" if photo.get("has_gps") else "sem GPS"
        doc.add_paragraph(f"{photo.get('filename')} | {display_dt(photo.get('taken_at')) or '-'} | {gps}")
    doc.add_heading("7. Pagamentos", level=1)
    for pay in pack["payments"]:
        doc.add_paragraph(
            f"{pay.get('kind')} | {pay.get('description') or '-'} | {pay.get('amount') or '-'} "
            f"{pay.get('currency') or ''} | {display_dt(pay.get('ts')) or '-'} | ID {pay.get('transaction_id') or '-'}"
        )
    doc.add_heading("8. Conexões encontradas", level=1)
    doc.add_paragraph("Cadeia: " + " → ".join(pack["graph"]["chain"]))
    for edge in pack["graph"]["edges"][:80]:
        if edge.get("label"):
            doc.add_paragraph(f"{edge['from']} —{edge['label']}→ {edge['to']}", style="List Bullet")
    doc.add_heading("9. Cadeia de custódia dos arquivos importados", level=1)
    for item in pack["imports"]:
        doc.add_paragraph(
            f"Arquivo original: {item.get('original_filename')}\n"
            f"SHA-256: {item.get('sha256')}\n"
            f"SHA-1: {item.get('sha1')}\n"
            f"MD5: {item.get('md5')}\n"
            f"Tamanho: {item.get('size_bytes')} bytes\n"
            f"Importado em: {display_dt(item.get('imported_at'))}"
        )
    doc.save(dest)
    return dest


def generate_pdf(case_id: int) -> Path:
    pack = _case_pack(case_id)
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    dest = REPORT_DIR / f"relatorio_forense_caso_{case_id}_{now_iso().replace(':', '')}.pdf"
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="Cover", fontSize=16, leading=20, textColor=colors.HexColor("#183a31"), spaceAfter=8, alignment=1))
    styles.add(ParagraphStyle(name="Section", fontSize=12, leading=16, textColor=colors.HexColor("#183a31"), spaceBefore=12, spaceAfter=6, fontName="Helvetica-Bold"))
    styles.add(ParagraphStyle(name="Body2", fontSize=9, leading=12, spaceAfter=4))
    doc = SimpleDocTemplate(str(dest), pagesize=A4, leftMargin=1.8 * cm, rightMargin=1.8 * cm, topMargin=1.6 * cm, bottomMargin=1.6 * cm)
    story = [
        Paragraph("POLÍCIA CIVIL DO PARANÁ", styles["Cover"]),
        Paragraph("ANALISADOR GOOGLE FORENSE — RELATÓRIO POLICIAL", styles["Cover"]),
        Paragraph(
            "Documento gerado automaticamente a partir de produção Google LERS. Conferir com os originais.",
            styles["Body2"],
        ),
        Paragraph("1. Identificação da conta", styles["Section"]),
    ]
    if not pack["accounts"]:
        story.append(Paragraph("Nenhuma conta extraída.", styles["Body2"]))
    for acc in pack["accounts"]:
        story.append(
            Paragraph(
                f"<b>{acc.get('display_name') or 'Conta'}</b><br/>"
                f"E-mail principal: {acc.get('primary_email') or '-'}<br/>"
                f"Alternativos: {', '.join(from_json(acc.get('alternate_emails'), []) or []) or '-'}<br/>"
                f"Telefones: {', '.join(from_json(acc.get('phones'), []) or []) or '-'}<br/>"
                f"Criação: {display_dt(acc.get('created_on')) or '-'} | "
                f"Última atividade: {display_dt(acc.get('last_activity')) or '-'}<br/>"
                f"Status: {acc.get('status') or '-'} | Exclusão: {display_dt(acc.get('deletion_date')) or '-'}",
                styles["Body2"],
            )
        )
    story.append(Paragraph("2. Dispositivos encontrados", styles["Section"]))
    rows = [["Modelo", "IMEI 1", "IMEI 2", "Série", "Contas"]]
    for device in pack["devices"]:
        rows.append(
            [
                device.get("model") or "-",
                device.get("imei1") or "-",
                device.get("imei2") or "-",
                device.get("serial_number") or "-",
                ", ".join(_emails_of_device(pack, device["id"])) or "-",
            ]
        )
    table = Table(rows, colWidths=[3.2 * cm, 3.5 * cm, 3.5 * cm, 3 * cm, 4.3 * cm])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#183a31")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#c5a253")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    story.append(table)
    if pack["alerts"]:
        story.append(Paragraph("Alerta: mesmo aparelho vinculado a múltiplas contas", styles["Section"]))
        for alert in pack["alerts"]:
            story.append(Paragraph(f"IMEI {alert['imei']}: {', '.join(alert['accounts'])}", styles["Body2"]))
    story.append(Paragraph("3. Linha temporal", styles["Section"]))
    for ev in pack["events"][:80]:
        story.append(
            Paragraph(
                f"{display_dt(ev.get('ts')) or '-'} — {ev.get('product') or '-'} — {ev.get('description') or '-'}",
                styles["Body2"],
            )
        )
    story.append(Paragraph("4. Locais relevantes", styles["Section"]))
    for loc in pack["locations"][:50]:
        story.append(
            Paragraph(
                f"{display_dt(loc.get('ts')) or '-'} | {loc.get('lat')}, {loc.get('lon')} | {loc.get('place_name') or '-'}",
                styles["Body2"],
            )
        )
    story.append(Paragraph("5. Arquivos importantes", styles["Section"]))
    important = [f for f in pack["files"] if from_json(f.get("keywords_hit"), [])]
    if not important:
        story.append(Paragraph("Nenhum arquivo com palavras-chave investigativas.", styles["Body2"]))
    for item in important:
        story.append(
            Paragraph(
                f"{item.get('name')} — {', '.join(from_json(item.get('keywords_hit'), []) or [])}",
                styles["Body2"],
            )
        )
    story.append(Paragraph("6. Fotos", styles["Section"]))
    for photo in pack["photos"][:40]:
        gps = f"{photo.get('lat')}, {photo.get('lon')}" if photo.get("has_gps") else "sem GPS"
        story.append(Paragraph(f"{photo.get('filename')} | {display_dt(photo.get('taken_at')) or '-'} | {gps}", styles["Body2"]))
    story.append(Paragraph("7. Pagamentos", styles["Section"]))
    for pay in pack["payments"]:
        story.append(
            Paragraph(
                f"{pay.get('kind')} | {pay.get('description') or '-'} | {pay.get('amount') or '-'} | ID {pay.get('transaction_id') or '-'}",
                styles["Body2"],
            )
        )
    story.append(Paragraph("8. Conexões encontradas", styles["Section"]))
    story.append(Paragraph(" → ".join(pack["graph"]["chain"]), styles["Body2"]))
    story.append(Paragraph("9. Cadeia de custódia", styles["Section"]))
    for item in pack["imports"]:
        story.append(
            Paragraph(
                f"{item.get('original_filename')}<br/>SHA-256: {item.get('sha256')}<br/>SHA-1: {item.get('sha1')}<br/>MD5: {item.get('md5')}",
                styles["Body2"],
            )
        )
    story.append(Spacer(1, 12))
    story.append(Paragraph(f"Gerado em {display_dt(now_iso())} — Analisador Google Forense", styles["Body2"]))
    doc.build(story)
    return dest
